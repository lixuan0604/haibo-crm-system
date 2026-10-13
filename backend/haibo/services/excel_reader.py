from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from config import IMPORT_YEAR

CENT = Decimal("0.01")
EXCEL_PATH = Path(__file__).resolve().parents[3] / "海博固定费用.xlsx"
PAYMENT_COLUMNS = (5, 6, 7, 8, 9)


def parse_workbook(path=None, year=None):
    path = Path(path or EXCEL_PATH)
    if not path.exists():
        raise FileNotFoundError(path)
    year = int(year or IMPORT_YEAR)
    workbook = load_workbook(path, data_only=False)
    worksheet = workbook["Sheet2"] if "Sheet2" in workbook.sheetnames else workbook.worksheets[0]
    ranges = _merge_ranges(worksheet)
    return {
        "year": year,
        "fixed": _parse_fixed(worksheet, ranges),
        "categories": _parse_categories(worksheet),
        "monthly": _parse_monthly(worksheet),
    }


def _merge_ranges(worksheet):
    ranges = []
    for merged in worksheet.merged_cells.ranges:
        value = worksheet.cell(merged.min_row, merged.min_col).value
        ranges.append((merged.min_row, merged.min_col, merged.max_row, merged.max_col, value))
    return ranges


def _anchor(ranges, row, col):
    for min_row, min_col, max_row, max_col, value in ranges:
        if min_col <= col <= max_col and min_row <= row <= max_row:
            return min_row, value
    return None, None


def _merged_value(worksheet, ranges, row, col):
    raw = worksheet.cell(row, col).value
    if raw is not None:
        return raw
    _row, value = _anchor(ranges, row, col)
    return value


def _clean_text(value):
    if value is None:
        return None
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return text or None


def _clean_name(value):
    text = _clean_text(value)
    if not text:
        return ""
    return " ".join(part.strip() for part in text.split("\n") if part.strip())


def _as_number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        text = value.strip().replace(",", "").replace("，", "")
        if not text or text.startswith("="):
            return None
        value = text
    try:
        if isinstance(value, float):
            value = f"{value:.2f}"
        return Decimal(str(value)).quantize(CENT)
    except (InvalidOperation, ValueError):
        return None


def _is_total(value):
    return isinstance(value, str) and "合计" in value


def _parse_fixed(worksheet, ranges):
    items = []
    for row in range(3, worksheet.max_row + 1):
        raw_name = worksheet.cell(row, 1).value
        raw_note = worksheet.cell(row, 2).value
        raw_period = worksheet.cell(row, 3).value
        if _is_total(raw_name) or _is_total(raw_period):
            break

        anchor_row, _value = _anchor(ranges, row, 1)
        is_child = raw_name is None and anchor_row is not None and anchor_row < row
        if raw_name is None and not is_child and raw_note is None:
            continue

        if is_child:
            name = _clean_name(raw_note)
            note = None
            period = _clean_text(raw_period)
            parent_row = anchor_row
        else:
            name = _clean_name(raw_name if raw_name is not None else raw_note)
            note = _clean_text(_merged_value(worksheet, ranges, row, 2)) if raw_name is not None else None
            period = _clean_text(_merged_value(worksheet, ranges, row, 3))
            parent_row = None
        if not name:
            continue

        annual = _as_number(worksheet.cell(row, 4).value)
        payments = []
        for seq_no, col in enumerate(PAYMENT_COLUMNS, start=1):
            amount = _as_number(worksheet.cell(row, col).value)
            if amount is None or amount == 0:
                continue
            payments.append({"seqNo": seq_no, "amount": abs(amount)})

        items.append(
            {
                "source_row": row,
                "parent_row": parent_row,
                "name": name,
                "note": note,
                "period_text": period,
                "annual_amount": annual if annual is not None else Decimal("0.00"),
                "sort_order": row,
                "payments": payments,
            }
        )
    return items


def _parse_categories(worksheet):
    categories = []
    for col in range(13, 22):
        header = _clean_text(worksheet.cell(2, col).value)
        if header and not header.startswith("="):
            categories.append(header)
    return categories


def _parse_month(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        month = int(value)
        return month if 1 <= month <= 12 else None
    text = str(value).strip()
    if not text or "合计" in text or text.startswith("="):
        return None
    digits = "".join(char for char in text if char.isdigit())
    if not digits:
        return None
    month = int(digits)
    return month if 1 <= month <= 12 else None


def _parse_monthly(worksheet):
    categories = _parse_categories(worksheet)
    cells = []
    for row in range(3, worksheet.max_row + 1):
        month = _parse_month(worksheet.cell(row, 12).value)
        if month is None:
            if _is_total(worksheet.cell(row, 12).value):
                break
            continue
        for offset, category in enumerate(categories):
            amount = _as_number(worksheet.cell(row, 13 + offset).value)
            if amount is None:
                continue
            cells.append({"month": month, "category": category, "amount": amount})
    return cells
